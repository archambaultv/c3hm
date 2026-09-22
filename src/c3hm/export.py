"""
Exports : tableur de notes pour Omnivox, grille en Markdown, rétroaction PDF.

Le rendu HTML de la rétroaction utilise les mêmes gabarits Jinja que l'interface
web, sans dépendre de Flask, pour que la ligne de commande fonctionne seule.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from jinja2 import Environment, PackageLoader, select_autoescape
from weasyprint import HTML

from c3hm.markup import render_markdown
from c3hm.model import (
    Evaluation,
    GradeResult,
    Notes,
    Student,
    compute_grade,
    level_color,
    merge_notes,
    round_half_up,
)
from c3hm.storage import Workspace

OUTPUT_DIR = "sortie"
FEEDBACK_DIR = "retroaction"
OMNIVOX_FILE = "notes_omnivox.xlsx"
RUBRIC_MD_FILE = "grille.md"
RUBRIC_PDF_FILE = "grille.pdf"
ZIP_FILE = "retroaction.zip"


class ExportError(Exception):
    pass


@dataclass
class StudentGrade:
    student: Student
    notes: Notes  # notes effectives (héritage d'équipe appliqué)
    own: Notes
    team_notes: Notes | None
    grade: GradeResult


def gather(workspace: Workspace, evaluation: Evaluation) -> list[StudentGrade]:
    """Notes effectives de chaque étudiant, équipes fusionnées."""
    cache: dict[str, Notes] = {}
    result: list[StudentGrade] = []
    for student in workspace.load_students():
        own = workspace.load_notes(student.matricule)
        team_notes = None
        if student.team:
            if student.team not in cache:
                cache[student.team] = workspace.load_team_notes(student.team)
            team_notes = cache[student.team]
        effective = merge_notes(team_notes, own) if team_notes is not None else own
        result.append(StudentGrade(student, effective, own, team_notes, compute_grade(evaluation, effective)))
    return result


def safe_filename(student: Student) -> str:
    parts = [student.surname, student.firstname, student.matricule]
    name = "_".join(p.strip() for p in parts if p.strip())
    return "".join("-" if c in '/\\:*?"<>|' else c for c in name)


# --- Markdown ----------------------------------------------------------------


def rubric_markdown(evaluation: Evaluation) -> str:
    """Grille en Markdown, pour un site Docusaurus. Tableau ou sections selon l'affichage choisi."""
    lines = [f"# {evaluation.title}", "", f"{evaluation.course} — {evaluation.session}", ""]
    if evaluation.layout == "liste":
        for crit in evaluation.criteria:
            lines += [f"## {crit.label} ({crit.weight} %)", ""]
            for lvl in evaluation.scale:
                text = crit.descriptors.get(lvl.label, "").strip()
                if text:
                    lines += [f"**{lvl.label} ({lvl.percentage} %)**", "", text, ""]
        return "\n".join(lines) + "\n"

    header = ["Critère"] + [f"{lvl.label} ({lvl.percentage} %)" for lvl in evaluation.scale]
    lines += [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for crit in evaluation.criteria:
        cells = [f"**{crit.label}** ({crit.weight} %)"]
        cells += [
            crit.descriptors.get(lvl.label, "").replace("\n", "<br>").replace("|", "\\|") for lvl in evaluation.scale
        ]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines) + "\n"


# --- Omnivox -----------------------------------------------------------------


def omnivox_rows(grades: list[StudentGrade]) -> str:
    """
    Les trois premières colonnes en texte séparé par des tabulations, à coller dans Omnivox.

    Les commentaires sont mis sur une seule ligne : Omnivox lit une note par ligne.
    """
    lines = []
    for item in grades:
        comment = " ".join(item.notes.comment.split())
        lines.append(f"{item.student.matricule}\t{item.grade.final}\t{comment}")
    return "\n".join(lines)


def write_omnivox_xlsx(grades: list[StudentGrade], path: Path) -> None:
    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notes pour Omnivox"
    ws.sheet_view.showGridLines = False
    ws.append(["Code omnivox", "Note", "Commentaire", "Nom"])
    for item in grades:
        ws.append([item.student.matricule, item.grade.final, item.notes.comment, item.student.fullname()])
    table = Table(displayName="NotesOmnivox", ref=f"A1:D{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    for column, width in (("A", 16), ("B", 10), ("C", 70), ("D", 34)):
        ws.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --- Rétroaction -------------------------------------------------------------


def _environment() -> Environment:
    env = Environment(
        loader=PackageLoader("c3hm.web", "templates"),
        autoescape=select_autoescape(["html"]),
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.globals.update(level_color=level_color, round_half_up=round_half_up)
    env.filters["md"] = render_markdown
    return env


def rubric_html(evaluation: Evaluation) -> str:
    return _environment().get_template("grille_pdf.html").render(evaluation=evaluation)


def feedback_html(evaluation: Evaluation, item: StudentGrade) -> str:
    return (
        _environment()
        .get_template("retroaction.html")
        .render(evaluation=evaluation, student=item.student, notes=item.notes, grade=item.grade)
    )


def write_pdf(html: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html).write_pdf(str(path))


def zip_files(files: list[Path], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for file in files:
            archive.write(file, file.name)


# --- Tout ensemble -----------------------------------------------------------


@dataclass
class ExportReport:
    output_dir: Path
    exported: list[str]  # étudiants inclus dans le tableur
    skipped: list[str]  # correction incomplète
    files: list[Path]
    pdf_count: int = 0


def clear_output(workspace: Workspace) -> Path:
    """Vide sortie/ au complet. Tout son contenu est produit par c3hm."""
    out = workspace.root / OUTPUT_DIR
    if out.is_dir():
        shutil.rmtree(out)
    out.mkdir(parents=True, exist_ok=True)
    return out


def clear_notes_output(workspace: Workspace) -> None:
    """
    Efface les fichiers de notes déjà exportés, sans toucher à la grille.

    Sinon la rétroaction d'un étudiant retiré de la liste resterait sur le disque.
    """
    out = workspace.root / OUTPUT_DIR
    shutil.rmtree(out / FEEDBACK_DIR, ignore_errors=True)
    for name in (OMNIVOX_FILE, ZIP_FILE):
        (out / name).unlink(missing_ok=True)


def _checked_evaluation(workspace: Workspace) -> Evaluation:
    evaluation = workspace.load_evaluation()
    errors = evaluation.validate()
    if errors:
        raise ExportError("L'évaluation est incomplète : " + errors[0])
    return evaluation


def export_rubric_markdown(workspace: Workspace) -> Path:
    evaluation = _checked_evaluation(workspace)
    path = workspace.root / OUTPUT_DIR / RUBRIC_MD_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(rubric_markdown(evaluation), encoding="utf-8")
    return path


def export_rubric_pdf(workspace: Workspace) -> Path:
    evaluation = _checked_evaluation(workspace)
    path = workspace.root / OUTPUT_DIR / RUBRIC_PDF_FILE
    write_pdf(rubric_html(evaluation), path)
    return path


def ready_grades(workspace: Workspace) -> list[StudentGrade]:
    """Étudiants dont la correction est complète, dans l'ordre alphabétique."""
    evaluation = _checked_evaluation(workspace)
    return [g for g in gather(workspace, evaluation) if g.grade.complete and g.grade.final is not None]


def export_notes(workspace: Workspace) -> ExportReport:
    """Tableur Omnivox, rétroactions PDF et archive zip. Efface d'abord ceux de l'export précédent."""
    evaluation = _checked_evaluation(workspace)
    clear_notes_output(workspace)
    out = workspace.root / OUTPUT_DIR
    grades = gather(workspace, evaluation)
    ready = [g for g in grades if g.grade.complete and g.grade.final is not None]
    skipped = [g.student.fullname() for g in grades if g not in ready]

    omnivox_path = out / OMNIVOX_FILE
    write_omnivox_xlsx(ready, omnivox_path)
    files: list[Path] = [omnivox_path]

    pdfs: list[Path] = []
    for item in ready:
        pdf_path = out / FEEDBACK_DIR / f"{safe_filename(item.student)}.pdf"
        write_pdf(feedback_html(evaluation, item), pdf_path)
        pdfs.append(pdf_path)
    if pdfs:
        zip_path = out / ZIP_FILE
        zip_files(pdfs, zip_path)
        files += [*pdfs, zip_path]

    return ExportReport(
        output_dir=out,
        exported=[g.student.fullname() for g in ready],
        skipped=skipped,
        files=files,
        pdf_count=len(pdfs),
    )


def export_all(workspace: Workspace) -> ExportReport:
    """Tout ce que produit c3hm : notes, rétroactions et grille vierge. Vide sortie/ d'abord."""
    _checked_evaluation(workspace)
    clear_output(workspace)
    report = export_notes(workspace)
    report.files += [export_rubric_markdown(workspace), export_rubric_pdf(workspace)]
    return report
